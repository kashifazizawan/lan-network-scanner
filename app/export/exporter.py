"""Export engine: CSV, XLSX, JSON, PDF, TXT."""
from __future__ import annotations

import csv
import json

from ..core.models import Device

COLUMNS = [
    "IP", "MAC", "Hostname", "Vendor", "Status",
    "Latency (ms)", "Open Ports", "Interface", "First Seen", "Last Seen",
]


def _rows(devices: list[Device]) -> list[list]:
    return [
        [
            d.ip, d.mac or "-", d.hostname, d.vendor, d.status,
            round(d.latency_ms, 1) if d.status == "Online" else "",
            d.ports_display, d.interface or "-", d.first_seen, d.last_seen,
        ]
        for d in devices
    ]


def export_csv(devices: list[Device], path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(COLUMNS)
        w.writerows(_rows(devices))


def export_json(devices: list[Device], path: str) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        json.dump([d.to_dict() for d in devices], fh, indent=2)


def export_txt(devices: list[Device], path: str) -> None:
    widths = [16, 18, 20, 20, 8, 10, 24, 12, 19, 19]
    lines = []
    header = "".join(f"{c:<{w}}" for c, w in zip(COLUMNS, widths)).rstrip()
    lines.append(header)
    lines.append("-" * len(header))
    for row in _rows(devices):
        lines.append("".join(f"{str(v):<{w}}" for v, w in zip(row, widths)).rstrip())
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")


def export_xlsx(devices: list[Device], path: str) -> None:
    """Requires openpyxl."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "Excel export needs the 'openpyxl' package. Install it with: pip install openpyxl"
        ) from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Scan Results"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for row in _rows(devices):
        ws.append(row)
    for col, width in zip("ABCDEFGHIJ", [16, 19, 22, 22, 9, 12, 26, 14, 20, 20]):
        ws.column_dimensions[col].width = width
    wb.save(path)


def export_pdf(devices: list[Device], path: str) -> None:
    """Requires reportlab."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError as exc:
        raise RuntimeError(
            "PDF export needs the 'reportlab' package. Install it with: pip install reportlab"
        ) from exc

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(path, pagesize=landscape(A4), title="LAN Scanner Report")
    story = [Paragraph(f"LAN Network Scan Report — {len(devices)} devices", styles["Title"]), Paragraph(" ", styles["Normal"])]
    data = [COLUMNS] + _rows(devices)
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
        ])
    )
    story.append(table)
    doc.build(story)


EXPORTERS = {
    "csv": export_csv,
    "json": export_json,
    "txt": export_txt,
    "xlsx": export_xlsx,
    "pdf": export_pdf,
}


def export(devices: list[Device], path: str, fmt: str) -> None:
    fmt = fmt.lower()
    if fmt not in EXPORTERS:
        raise ValueError(f"Unsupported export format: '{fmt}'. Use csv, xlsx, json, pdf or txt.")
    EXPORTERS[fmt](devices, path)
