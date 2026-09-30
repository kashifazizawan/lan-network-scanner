"""Tests for the export engine (CSV, JSON, TXT, XLSX; PDF optional)."""
import json
import csv

import pytest

from app.core.models import Device
from app.export.exporter import export, export_csv, export_json, export_txt, export_xlsx


@pytest.fixture
def devices():
    return [
        Device(ip="192.168.1.1", mac="AA-BB-CC-DD-EE-FF", hostname="Router",
               vendor="TP-Link", status="Online", latency_ms=2.0, open_ports=[80, 443]),
        Device(ip="192.168.1.20", mac="11-22-33-44-55-66", hostname="OFFICE-PC",
               vendor="Dell Inc.", status="Online", latency_ms=4.0, open_ports=[135, 445]),
    ]


def _read(path):
    with open(path, encoding="utf-8-sig" if path.endswith(".csv") else "utf-8") as fh:
        return fh.read()


class TestCsv:
    def test_header_and_rows(self, tmp_path, devices):
        p = str(tmp_path / "out.csv")
        export_csv(devices, p)
        text = _read(p)
        assert text.splitlines()[0] == "IP,MAC,Hostname,Vendor,Status,Latency (ms),Open Ports,Interface,First Seen,Last Seen"
        assert "192.168.1.1" in text
        assert "80,443" in text
        assert "OFFICE-PC" in text


class TestJson:
    def test_roundtrip(self, tmp_path, devices):
        p = str(tmp_path / "out.json")
        export_json(devices, p)
        data = json.loads(_read(p))
        assert len(data) == 2
        assert data[0]["ip"] == "192.168.1.1"
        assert data[0]["open_ports"] == [80, 443]

    def test_device_roundtrip(self, devices):
        d = devices[0].to_dict()
        d2 = Device.from_dict(d)
        assert d2.ip == d["ip"]
        assert d2.mac == d["mac"]
        assert d2.open_ports == [80, 443]


class TestTxt:
    def test_columns_aligned(self, tmp_path, devices):
        p = str(tmp_path / "out.txt")
        export_txt(devices, p)
        lines = _read(p).splitlines()
        assert "IP" in lines[0]
        assert "192.168.1.1" in lines[2]


class TestXlsx:
    def test_xlsx(self, tmp_path, devices):
        openpyxl = pytest.importorskip("openpyxl")
        p = str(tmp_path / "out.xlsx")
        export_xlsx(devices, p)
        from openpyxl import load_workbook
        wb = load_workbook(p)
        ws = wb.active
        assert ws.cell(1, 1).value == "IP"
        assert ws.cell(2, 1).value == "192.168.1.1"


class TestDispatch:
    def test_export_by_format(self, tmp_path, devices):
        for fmt in ("csv", "json", "txt"):
            p = str(tmp_path / f"out.{fmt}")
            export(devices, p, fmt)
            assert _read(p)

    def test_unknown_format(self, tmp_path, devices):
        with pytest.raises(ValueError):
            export(devices, str(tmp_path / "out.xml"), "xml")
